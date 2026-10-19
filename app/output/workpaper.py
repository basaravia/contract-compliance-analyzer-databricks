"""Papel de trabajo de auditoría en Excel (Hoja 1 «Papel de trabajo» + Hoja 2 «Anexo técnico»).

Formato de papel de trabajo del comparador normativo,
adaptado a «reglas vs. contrato». 100 % Python con openpyxl: **cero llamadas al modelo**; este
módulo no importa nada de `app.llm`. El modelo produce campos, el código produce el formato.

Dos decisiones deterministas viven aquí, a la vista y con pruebas:

1. `classify_search`: convierte (estado del evaluador, evidencia, estado de extracción) en un
   resultado de búsqueda explícito: *encontrado*, *no_encontrado*, *revision* o *no_determinable*.
   «No se encontró en el contrato» es un resultado propio, no un error ni un «no cumple» implícito.
2. `verify_literal`: toda evidencia citada debe ser subcadena del texto extraído del contrato.
   Si no lo es, se marca `cita_no_verificada` (nunca se oculta).
"""
from __future__ import annotations

import io
import math
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from datetime import date
from typing import Literal

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SearchOutcome = Literal["encontrado", "no_encontrado", "revision", "no_determinable"]
Mark = Literal["A", "L", "R", "N", "V"]

# --- Marcas (docs/arquitectura/04-papel-de-trabajo.md) -------------------------------------
MARK_MEANING = {
    "A": "Sí cumple",
    "L": "Parcialmente",
    "R": "No cumple (la cláusula existe pero no satisface la regla)",
    "N": "No se encontró en el contrato",
    "V": "Requiere revisión del auditor (no determinable o salida no confiable)",
}
MARK_COLOR = {"A": "C6EFCE", "L": "FFEB9C", "R": "F8CBAD", "N": "FCE4D6", "V": "DDEBF7"}
PRIMARY = "1F6F4A"  # verde institucional (brand kit pendiente, igual que en v2)

MAX_CELL = 32767  # límite de Excel por celda
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")  # CWE-1236 / OWASP: Excel los evalúa como fórmula
COLUMNS = ["Regla / Referencia normativa", "Requisito a validar", "Contrato (evidencia literal)",
           "Verificación", "Comentario", "Referencias / Evidencias"]
WIDTHS = [30, 45, 50, 12, 45, 32]
HEADER_LABELS = ["Nombre de la revisión", "Nombre del papel de trabajo", "Objetivo", "Corte o periodo",
                 "Fuente", "Marcas de verificación", "Conclusiones", "Elaborado por", "Revisado por",
                 "Fecha de ejecución"]
OBJECTIVE = ("Evaluar el cumplimiento del contrato frente al conjunto de reglas normativas definidas, "
             "buscando en el contrato la evidencia de cada regla y dejando constancia expresa de las "
             "reglas cuya cláusula no se encontró.")
DRAFT_PREFIX = "[Borrador generado automáticamente — validar]"
DISCLAIMER = "Herramienta de apoyo; no constituye asesoría legal. Debe revisarla un profesional del derecho."


# --- Modelo de datos de entrada (sin dependencia de SQLAlchemy: se puede probar sola) ------
@dataclass
class WorkpaperRow:
    rule_id: str
    rule_name: str
    rule_description: str
    category: str | None
    severity: str
    citations: list[str]
    status: str  # estado interno del evaluador
    justification: str
    evidence_text: str | None
    page: int | None
    bbox: tuple[float, float, float, float] | None
    confidence: float
    edited: bool = False
    is_manual: bool = False
    # Rellenados por `enrich`:
    outcome: SearchOutcome = "revision"
    mark: Mark = "V"
    literal_verified: bool | None = None  # None = no aplica (no hay evidencia que verificar)
    flags: list[str] = field(default_factory=list)
    # De la traza (si existe):
    retrieved: int | None = None
    best_score: float | None = None
    retrieved_pages: list[int] = field(default_factory=list)
    second_pass: str | None = None  # found | confirmed_absent | partial_scan | failed | None
    chunks_scanned: int | None = None
    chunks_total: int | None = None


@dataclass
class WorkpaperInput:
    contract_name: str
    contract_pages: int
    extraction_status: str
    rows: list[WorkpaperRow]
    analysis_id: str
    analysis_trace: dict | None = None
    execution_date: date | None = None


# --- Decisiones deterministas ----------------------------------------------------------------
def classify_search(status: str, evidence_text: str | None, extraction_status: str) -> tuple[SearchOutcome, Mark]:
    """Resultado de la búsqueda y marca. Reglas, en orden:

    1. Contrato sin texto utilizable (escaneado sin OCR / extracción fallida) → no se pudo buscar.
    2. `requiere_revision` → revisión humana (salida del evaluador no confiable o incierta).
    3. Hay evidencia → *encontrado*; la marca sale del estado (cumple A, parcial L, no_cumple R).
       Un `cumple`/`parcial` **sin** evidencia es inconsistente → revisión (no se acepta sin cita).
    4. `no_cumple` o `no_determinable` sin evidencia → **no encontrado** (marca N).
    """
    if extraction_status != "done":
        return "no_determinable", "V"
    has_evidence = bool(evidence_text and evidence_text.strip())
    if status == "requiere_revision":
        return "revision", "V"
    if has_evidence:
        return "encontrado", {"cumple": "A", "parcial": "L", "no_cumple": "R"}.get(status, "V")
    if status in ("no_cumple", "no_determinable"):
        return "no_encontrado", "N"
    return "revision", "V"  # cumple/parcial sin evidencia


def _norm(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    return re.sub(r"\s+", " ", text).strip().casefold()


def verify_literal(evidence: str | None, page_texts: dict[int, str], page: int | None) -> bool | None:
    """¿La evidencia es subcadena literal del texto del contrato? (espacios y mayúsculas normalizados).

    Si hay página, se busca primero ahí; si no aparece, en todo el documento (los fragmentos pueden
    cruzar el salto de página). `None` si no hay evidencia que verificar."""
    if not evidence or not evidence.strip():
        return None
    needle = _norm(evidence)
    if page is not None and page in page_texts and needle in _norm(page_texts[page]):
        return True
    return needle in _norm(" ".join(page_texts[p] for p in sorted(page_texts)))


def enrich(inp: WorkpaperInput, page_texts: dict[int, str]) -> WorkpaperInput:
    """Calcula resultado de búsqueda, marca, verificación literal y banderas de cada fila."""
    for r in inp.rows:
        r.outcome, r.mark = classify_search(r.status, r.evidence_text, inp.extraction_status)
        r.literal_verified = verify_literal(r.evidence_text, page_texts, r.page)
        r.flags = []
        if r.literal_verified is False:
            r.flags.append("cita_no_verificada")
        if r.outcome == "revision":
            r.flags.append("requiere_revision")
        if r.status in ("cumple", "parcial") and not (r.evidence_text or "").strip():
            r.flags.append("sin_evidencia")
        if r.second_pass == "partial_scan":
            r.flags.append("busqueda_parcial")
        if r.second_pass == "failed":
            r.flags.append("segunda_pasada_fallida")
        if r.edited:
            r.flags.append("editado_manualmente")
        if r.is_manual:
            r.flags.append("hallazgo_manual")
        # Una cita que no se pudo verificar nunca deja una marca A/L/R sin advertencia: pasa a revisión.
        if r.literal_verified is False and r.mark in ("A", "L", "R"):
            r.mark = "V"
            r.outcome = "revision"
    return inp


def mark_counts(rows: list[WorkpaperRow]) -> dict[str, int]:
    c = Counter(r.mark for r in rows)
    return {m: c.get(m, 0) for m in "ALRNV"}


def draft_conclusion(inp: WorkpaperInput) -> str:
    """Conclusión determinista (sin modelo): cifras y reglas a atender, nada inventado."""
    n = len(inp.rows)
    c = mark_counts(inp.rows)
    partes = [
        f"Se evaluaron {n} regla(s) contra el contrato «{inp.contract_name}»: "
        f"{c['A']} cumple(n), {c['L']} parcial(es), {c['R']} no cumple(n), "
        f"{c['N']} sin cláusula encontrada y {c['V']} por revisar."
    ]
    for etiqueta, marca in (("No cumplen", "R"), ("Parciales", "L"), ("No encontradas en el contrato", "N"),
                            ("Por revisar", "V")):
        nombres = [r.rule_name for r in inp.rows if r.mark == marca]
        if nombres:
            partes.append(f"{etiqueta}: " + "; ".join(nombres) + ".")
    partes.append(DISCLAIMER)
    return "\n".join(partes)


# --- Escritura segura de celdas --------------------------------------------------------------
def safe_text(value, page: int | None = None) -> str:
    """Texto listo para una celda: sin caracteres ilegales y con tope de 32.767 caracteres."""
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    if len(text) <= MAX_CELL:
        return text
    notice = f"[texto truncado, ver página {page}]" if page else "[texto truncado]"
    return text[: MAX_CELL - len(notice) - 1] + "\n" + notice


_thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="top", horizontal="center")


def put(ws, row: int, col: int, value, *, page: int | None = None, font: Font | None = None,
        fill: PatternFill | None = None, align: Alignment = TOP):
    """Única puerta de entrada de texto a las celdas. Números como número; el resto como TEXTO.

    Un texto que empieza por `= + - @` (o tab/CR) se fuerza a cadena con `quotePrefix` para que
    Excel nunca lo evalúe como fórmula (el contrato y la salida del modelo son entrada no confiable)."""
    cell = ws.cell(row=row, column=col)
    if isinstance(value, float) and not math.isfinite(value):
        cell.value = None
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        cell.value = value
    else:
        text = safe_text("" if value is None else value, page)
        cell.value = text
        if text.startswith(FORMULA_PREFIXES):
            cell.data_type = "s"
            cell.quotePrefix = True
    cell.alignment = align
    cell.border = BORDER
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    return cell


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", start_color=color, end_color=color)


HEADER_FONT = Font(bold=True, size=9, color="FFFFFF")
LABEL_FONT = Font(bold=True, size=9)


def _row_height(pairs: list[tuple[str, int]]) -> float:
    """Alto aproximado según el texto de cada celda (tope de Excel: 409 pt)."""
    lines = 1
    for text, width in pairs:
        n = sum(max(1, math.ceil(len(part) / max(1, width * 1.1))) for part in str(text).split("\n"))
        lines = max(lines, n)
    return min(409.0, 13.0 * lines + 4)


# --- Hoja 1 ----------------------------------------------------------------------------------
def _evidence_cell(r: WorkpaperRow) -> str:
    if r.outcome == "no_encontrado":
        return "— No se encontró en el contrato una cláusula que atienda esta regla —"
    if r.outcome == "no_determinable":
        return "— No se pudo buscar: el contrato no tiene texto extraíble —"
    if not (r.evidence_text or "").strip():
        return ""
    return r.evidence_text.strip() + (f"\n(p. {r.page})" if r.page else "")


def _refs_cell(r: WorkpaperRow, contract: str) -> str:
    lines = []
    if r.page and r.outcome == "encontrado":
        lines.append(f"Contrato: {contract}, p. {r.page}")
    if r.retrieved is not None:
        score = f", mejor similitud {r.best_score:.3f}" if r.best_score is not None else ""
        lines.append(f"Búsqueda: {r.retrieved} fragmento(s) revisado(s){score}")
    if r.second_pass == "found":
        lines.append(f"Hallado al ampliar la búsqueda a {r.chunks_scanned} de {r.chunks_total} fragmentos")
    elif r.second_pass == "confirmed_absent":
        lines.append(f"Ausencia confirmada: se revisó todo el contrato ({r.chunks_scanned} de {r.chunks_total} fragmentos)")
    elif r.second_pass == "partial_scan":
        lines.append(f"Aviso: búsqueda NO exhaustiva ({r.chunks_scanned} de {r.chunks_total} fragmentos); validar manualmente")
    elif r.second_pass == "failed":
        lines.append("Aviso: la búsqueda ampliada falló; no se pudo confirmar la ausencia")
    lines.append(f"Confianza del evaluador: {r.confidence:.2f}")
    if "cita_no_verificada" in r.flags:
        lines.append("Aviso: la cita no se verificó contra el texto del contrato")
    if "requiere_revision" in r.flags:
        lines.append("Requiere revisión del auditor")
    if "editado_manualmente" in r.flags:
        lines.append("Hallazgo editado manualmente")
    return "\n".join(lines)


def _sheet1(ws, inp: WorkpaperInput) -> None:
    conclusion = draft_conclusion(inp)
    legend = "  ·  ".join(f"{m} = {MARK_MEANING[m]}" for m in "ALRNV")
    refs = sorted({c for r in inp.rows for c in r.citations})
    source = (f"{len(inp.rows)} regla(s)" + (": " + "; ".join(refs) if refs else "")
              + f"\nversus\n{inp.contract_name} ({inp.contract_pages} pág.)")
    values = ["", f"Evaluación de cumplimiento contractual — {inp.contract_name}", OBJECTIVE, "", source, legend,
              f"{DRAFT_PREFIX}\n{conclusion}", "", "", (inp.execution_date or date.today()).strftime("%d/%m/%Y")]
    value_width = sum(WIDTHS[1:])
    for i, (label, value) in enumerate(zip(HEADER_LABELS, values), start=1):
        put(ws, i, 1, label, font=LABEL_FONT, fill=_fill("E2EFDA"))
        put(ws, i, 2, value)
        for col in range(3, 7):
            ws.cell(row=i, column=col).border = BORDER
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
        ws.row_dimensions[i].height = _row_height([(value, value_width), (label, WIDTHS[0])])

    h = len(HEADER_LABELS) + 2
    for col, name in enumerate(COLUMNS, start=1):
        put(ws, h, col, name, font=HEADER_FONT, fill=_fill(PRIMARY), align=CENTER)
        ws.column_dimensions[chr(64 + col)].width = WIDTHS[col - 1]
    ws.freeze_panes = ws.cell(row=h + 1, column=1)

    for n, r in enumerate(inp.rows, start=h + 1):
        cells = [
            r.rule_name + ("\n" + "; ".join(r.citations) if r.citations else ""),
            r.rule_description,
            _evidence_cell(r),
            r.mark,
            r.justification,
            _refs_cell(r, inp.contract_name),
        ]
        for col, text in enumerate(cells, start=1):
            if col == 4:
                put(ws, n, col, text, font=Font(bold=True), fill=_fill(MARK_COLOR[r.mark]), align=CENTER)
            else:
                put(ws, n, col, text, page=r.page)
        ws.row_dimensions[n].height = _row_height([(cells[i], WIDTHS[i]) for i in (0, 1, 2, 4, 5)])


# --- Hoja 2 ----------------------------------------------------------------------------------
ANNEX_COLUMNS = ["ID regla", "Regla", "Categoría", "Severidad", "Estado del evaluador", "Resultado de la búsqueda",
                 "Marca", "Fragmentos revisados", "Mejor similitud", "Páginas revisadas", "Página de la evidencia",
                 "Cita verificada", "Confianza", "Banderas", "Segunda pasada", "Fragmentos escaneados / total"]
ANNEX_WIDTHS = [38, 34, 16, 11, 18, 20, 8, 12, 12, 14, 12, 12, 10, 34, 18, 18]
SECOND_PASS_LABEL = {"found": "Hallado", "confirmed_absent": "Ausencia confirmada",
                     "partial_scan": "Búsqueda parcial", "failed": "Falló"}
OUTCOME_LABEL = {"encontrado": "Encontrado", "no_encontrado": "No encontrado", "revision": "Requiere revisión",
                 "no_determinable": "No determinable"}


def _sheet2(ws, inp: WorkpaperInput) -> None:
    put(ws, 1, 1, "Anexo técnico — todas las reglas evaluadas (no hay selección sesgada de evidencia)",
        font=Font(bold=True, size=11))
    t = inp.analysis_trace or {}
    prov = (f"Embeddings: {t.get('embedding_provider', 'n/d')} · LLM: {t.get('llm_provider', 'n/d')}"
            f"{' (' + t['llm_model'] + ')' if t.get('llm_model') else ''} · top_k: {t.get('retrieval_top_k', 'n/d')}"
            f" · análisis {inp.analysis_id}")
    put(ws, 2, 1, prov)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(ANNEX_COLUMNS))
    if t.get("limitations"):
        put(ws, 3, 1, "Limitaciones: " + " ".join(t["limitations"]))
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(ANNEX_COLUMNS))
    h = 5
    for col, name in enumerate(ANNEX_COLUMNS, start=1):
        put(ws, h, col, name, font=HEADER_FONT, fill=_fill(PRIMARY), align=CENTER)
        ws.column_dimensions[chr(64 + col)].width = ANNEX_WIDTHS[col - 1]
    ws.freeze_panes = ws.cell(row=h + 1, column=1)
    n = h + 1
    for r in inp.rows:
        verified = {True: "Sí", False: "No", None: "—"}[r.literal_verified]
        vals = [r.rule_id, r.rule_name, r.category or "", r.severity, r.status, OUTCOME_LABEL[r.outcome], r.mark,
                r.retrieved if r.retrieved is not None else "", round(r.best_score, 3) if r.best_score is not None else "",
                ", ".join(map(str, r.retrieved_pages)), r.page if r.page else "", verified, round(r.confidence, 2),
                ", ".join(r.flags), SECOND_PASS_LABEL.get(r.second_pass or "", "—"),
                f"{r.chunks_scanned} / {r.chunks_total}" if r.second_pass else "—"]
        for col, v in enumerate(vals, start=1):
            put(ws, n, col, v, fill=_fill(MARK_COLOR[r.mark]) if col == 7 else None, align=CENTER if col == 7 else TOP)
        n += 1

    n += 1
    for title, outcome in (("Reglas sin cláusula encontrada en el contrato", "no_encontrado"),
                           ("Reglas que no se pudieron buscar (contrato sin texto extraíble)", "no_determinable"),
                           ("Reglas que requieren revisión del auditor", "revision")):
        group = [r for r in inp.rows if r.outcome == outcome]
        put(ws, n, 1, f"{title}: {len(group)}", font=Font(bold=True))
        n += 1
        for r in group:
            put(ws, n, 1, r.rule_id)
            put(ws, n, 2, r.rule_name)
            put(ws, n, 3, r.justification)
            ws.merge_cells(start_row=n, start_column=3, end_row=n, end_column=len(ANNEX_COLUMNS))
            n += 1
        n += 1

    put(ws, n, 1, "Conteo por marca (debe coincidir con la Hoja 1)", font=Font(bold=True))
    n += 1
    for m, v in mark_counts(inp.rows).items():
        put(ws, n, 1, m)
        put(ws, n, 2, MARK_MEANING[m])
        put(ws, n, 3, v)
        n += 1


# --- API pública -----------------------------------------------------------------------------
def build_workpaper(inp: WorkpaperInput) -> bytes:
    """Genera el .xlsx en memoria. `inp` ya debe estar enriquecido (`enrich`)."""
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Papel de trabajo"
    _sheet1(ws1, inp)
    _sheet2(wb.create_sheet("Anexo técnico"), inp)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
