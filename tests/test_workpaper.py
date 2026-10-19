"""Pruebas unitarias del papel de trabajo: decisiones deterministas y formato Excel (sin BD, sin red)."""
from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from app.output.workpaper import (
    COLUMNS,
    WorkpaperInput,
    WorkpaperRow,
    build_workpaper,
    classify_search,
    draft_conclusion,
    enrich,
    mark_counts,
    verify_literal,
)

PAGES = {1: "El Proveedor  se compromete a tratar los\ndatos personales conforme a la ley.", 2: "Texto de la página dos."}


def row(rule_id="r1", status="cumple", evidence="se compromete a tratar los datos personales", page=1, **kw):
    return WorkpaperRow(
        rule_id=rule_id, rule_name=f"Regla {rule_id}", rule_description="Descripción", category="c",
        severity="alta", citations=["LOPDP Art. 5"], status=status, justification="Porque sí.",
        evidence_text=evidence, page=page, bbox=None, confidence=0.9, **kw,
    )


@pytest.mark.parametrize(
    "status,evidence,extraction,expected",
    [
        ("cumple", "texto", "done", ("encontrado", "A")),
        ("parcial", "texto", "done", ("encontrado", "L")),
        ("no_cumple", "texto", "done", ("encontrado", "R")),
        ("no_cumple", None, "done", ("no_encontrado", "N")),
        ("no_determinable", None, "done", ("no_encontrado", "N")),
        ("no_determinable", "  ", "done", ("no_encontrado", "N")),
        ("requiere_revision", None, "done", ("revision", "V")),
        ("requiere_revision", "texto", "done", ("revision", "V")),
        ("cumple", None, "done", ("revision", "V")),  # cumple sin cita no se acepta
        ("parcial", None, "done", ("revision", "V")),
        ("no_determinable", None, "scanned_no_ocr", ("no_determinable", "V")),  # no se pudo buscar
        ("cumple", "texto", "failed", ("no_determinable", "V")),
    ],
)
def test_classify_search(status, evidence, extraction, expected):
    assert classify_search(status, evidence, extraction) == expected


def test_verify_literal_normalizes_whitespace_and_case():
    assert verify_literal("SE COMPROMETE a tratar   los datos personales", PAGES, 1) is True


def test_verify_literal_detects_invented_text():
    assert verify_literal("El proveedor pagará una multa de USD 1.000", PAGES, 1) is False


def test_verify_literal_falls_back_to_whole_document_and_handles_no_evidence():
    assert verify_literal("Texto de la página dos.", PAGES, 1) is True  # página mal indicada, texto real
    assert verify_literal(None, PAGES, None) is None
    assert verify_literal("   ", PAGES, None) is None


def test_enrich_downgrades_unverified_citation_to_review():
    inp = WorkpaperInput("c.pdf", 2, "done", [row(evidence="texto inventado que no está en el contrato")], "a1")
    r = enrich(inp, PAGES).rows[0]
    assert r.mark == "V" and r.outcome == "revision" and "cita_no_verificada" in r.flags


def _full_input():
    rows = [
        row("a", "cumple"),
        row("l", "parcial"),
        row("r", "no_cumple"),
        row("n", "no_cumple", evidence=None, page=None),
        row("v", "requiere_revision", evidence=None, page=None),
    ]
    return enrich(WorkpaperInput("contrato.pdf", 2, "done", rows, "an-1", {"llm_provider": "mock", "retrieval_top_k": 5}), PAGES)


def test_mark_counts_and_conclusion_are_deterministic():
    inp = _full_input()
    assert mark_counts(inp.rows) == {"A": 1, "L": 1, "R": 1, "N": 1, "V": 1}
    text = draft_conclusion(inp)
    assert "5 regla(s)" in text and "Regla n" in text and "no constituye asesoría legal" in text


def test_excel_structure_matches_spec():
    inp = _full_input()
    wb = load_workbook(io.BytesIO(build_workpaper(inp)))
    assert wb.sheetnames == ["Papel de trabajo", "Anexo técnico"]
    ws1 = wb["Papel de trabajo"]
    header_row = 12
    assert [ws1.cell(header_row, c).value for c in range(1, 7)] == COLUMNS
    assert ws1.freeze_panes == f"A{header_row + 1}"
    assert ws1["A1"].value == "Nombre de la revisión" and ws1["A7"].value == "Conclusiones"
    assert "[Borrador generado automáticamente" in ws1["B7"].value
    marks = [ws1.cell(r, 4).value for r in range(header_row + 1, header_row + 1 + len(inp.rows))]
    assert marks == ["A", "L", "R", "N", "V"]
    # N dice explícitamente que no se encontró, y no inventa evidencia
    assert "No se encontró en el contrato" in ws1.cell(header_row + 4, 3).value


def test_sheet1_counts_equal_sheet2_counts():
    inp = _full_input()
    wb = load_workbook(io.BytesIO(build_workpaper(inp)))
    ws1, ws2 = wb["Papel de trabajo"], wb["Anexo técnico"]
    from_sheet1 = {m: 0 for m in "ALRNV"}
    for r in range(13, 13 + len(inp.rows)):
        from_sheet1[ws1.cell(r, 4).value] += 1
    start = next(r for r in range(1, ws2.max_row + 1) if str(ws2.cell(r, 1).value).startswith("Conteo por marca"))
    from_sheet2 = {ws2.cell(r, 1).value: ws2.cell(r, 3).value for r in range(start + 1, start + 6)}
    assert from_sheet1 == from_sheet2 == mark_counts(inp.rows)


@pytest.mark.parametrize("payload", ["=HYPERLINK(\"http://x\",\"a\")", "+cmd|' /C calc'!A0", "-2+3", "@SUM(A1)"])
def test_formula_injection_is_neutralised(payload):
    r = row(evidence=payload)
    r.justification = payload
    inp = enrich(WorkpaperInput("c.pdf", 1, "done", [r], "a"), {1: payload})
    ws = load_workbook(io.BytesIO(build_workpaper(inp)))["Papel de trabajo"]
    for col in (3, 5):
        cell = ws.cell(13, col)
        assert cell.value.startswith(payload)  # el texto no se altera (la evidencia añade "(p. N)")
        assert cell.data_type == "s"  # se guarda como cadena, no como fórmula
        assert cell.quotePrefix is True  # Excel no lo evalúa ni al editar la celda


def test_long_text_is_truncated_with_notice_and_illegal_chars_removed():
    r = row(evidence="x" * 40_000 + "\x00\x01")
    inp = enrich(WorkpaperInput("c.pdf", 1, "done", [r], "a"), {1: "x" * 40_000})
    ws = load_workbook(io.BytesIO(build_workpaper(inp)))["Papel de trabajo"]
    value = ws.cell(13, 3).value
    assert len(value) <= 32767 and "[texto truncado" in value and "\x00" not in value
